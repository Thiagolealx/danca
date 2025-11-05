# relatorios.py
from django.http import HttpResponse
from django.utils.timezone import now
from django.shortcuts import get_object_or_404
from django.db.models import Sum, Q
from django.db.models.functions import Lower
from django.contrib.contenttypes.models import ContentType
from django.db.models import Case, When, Value, DecimalField, F
from django.db.models.functions import Coalesce
from django.views import View
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import Entrada, Inscricao, Evento, Profissional, Planejamento, PedidoCamisa, Pagamento, InscricaoEvento, ProfissionalEvento, BaileAvulso, ParticipanteBaile, Saida
from django.shortcuts import render
import csv
from datetime import datetime
from django.template.loader import render_to_string
from .views import InscricaoListView
from collections import Counter

class InscricaoRelatorioDocxView(InscricaoListView):
    """ Gera relatório .docx com tabelas organizadas """
    
    def get(self, request, *args, **kwargs):
        # Remove a paginação para pegar todos os registros
        self.paginate_by = None
        queryset = self.get_queryset()

        document = Document()
        
        # Cabeçalho
        document.add_heading('Relatório de Inscrições', 0)
        
        # Informações gerais
        info_paragraph = document.add_paragraph()
        info_paragraph.add_run('Gerado em: ').bold = True
        info_paragraph.add_run(now().strftime("%d/%m/%Y %H:%M"))
        
        info_paragraph = document.add_paragraph()
        info_paragraph.add_run('Total de inscrições: ').bold = True
        info_paragraph.add_run(str(queryset.count()))
        
        document.add_paragraph()
        
        # Contadores
        uf_counter = Counter()
        status_counter = Counter()
        categoria_counter = Counter()

        # Função para tratar valores nulos
        def safe_str(value):
            if value is None:
                return "Não informado"
            return str(value)

        # Dados para tabela principal
        inscricoes_data = []
        for inscricao in queryset:
            categoria = safe_str(inscricao.categoria.descricao if inscricao.categoria else 'Sem categoria')
            
            if inscricao.valor_restante_db <= 0:
                status = 'Pago'
            elif inscricao.valor_pago_db > 0:
                status = 'Parcial'
            else:
                status = 'Pendente'
            
            uf = safe_str(getattr(inscricao, 'uf', 'Não informado'))
            
            # Atualizar contadores
            uf_counter[uf] += 1
            status_counter[status] += 1
            categoria_counter[categoria] += 1
            
            inscricoes_data.append({
                'nome': safe_str(inscricao.nome),
                'cpf': safe_str(inscricao.cpf),
                'categoria': categoria,
                'status': status,
                'uf': uf
            })

        # Tabela principal de inscrições
        document.add_heading('Lista de Inscrições', level=1)
        
        if inscricoes_data:
            table = document.add_table(rows=1, cols=5)
            table.style = 'Light Grid Accent 1'
            
            # Cabeçalho
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '#'
            hdr_cells[1].text = 'Nome'
            hdr_cells[2].text = 'CPF'
            hdr_cells[3].text = 'Categoria'
            hdr_cells[4].text = 'Status/UF'
            
            # Dados
            for i, inscricao in enumerate(inscricoes_data, 1):
                row_cells = table.add_row().cells
                row_cells[0].text = safe_str(i)
                row_cells[1].text = inscricao['nome']
                row_cells[2].text = inscricao['cpf']
                row_cells[3].text = inscricao['categoria']
                row_cells[4].text = f"{inscricao['status']} - {inscricao['uf']}"

        document.add_paragraph()

        # Resumo por Status
        document.add_heading('Resumo por Status', level=1)
        
        if status_counter:
            status_table = document.add_table(rows=1, cols=2)
            status_table.style = 'Light Grid Accent 1'
            
            hdr_cells = status_table.rows[0].cells
            hdr_cells[0].text = 'Status'
            hdr_cells[1].text = 'Quantidade'
            
            for status, quantidade in sorted(status_counter.items()):
                row_cells = status_table.add_row().cells
                row_cells[0].text = safe_str(status)
                row_cells[1].text = safe_str(quantidade)

        document.add_paragraph()

        # Resumo por UF
        document.add_heading('Resumo por UF', level=1)
        
        if uf_counter:
            uf_table = document.add_table(rows=1, cols=2)
            uf_table.style = 'Light Grid Accent 1'
            
            hdr_cells = uf_table.rows[0].cells
            hdr_cells[0].text = 'UF'
            hdr_cells[1].text = 'Quantidade'
            
            for uf, quantidade in sorted(uf_counter.items()):
                row_cells = uf_table.add_row().cells
                row_cells[0].text = safe_str(uf)
                row_cells[1].text = safe_str(quantidade)

        document.add_paragraph()

        # Resumo por Categoria
        document.add_heading('Resumo por Categoria', level=1)
        
        if categoria_counter:
            cat_table = document.add_table(rows=1, cols=2)
            cat_table.style = 'Light Grid Accent 1'
            
            hdr_cells = cat_table.rows[0].cells
            hdr_cells[0].text = 'Categoria'
            hdr_cells[1].text = 'Quantidade'
            
            for categoria, quantidade in sorted(categoria_counter.items()):
                row_cells = cat_table.add_row().cells
                row_cells[0].text = safe_str(categoria)
                row_cells[1].text = safe_str(quantidade)

        # Preparar resposta
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="relatorio_inscricoes.docx"'
        document.save(response)
        return response

def evento_inscritos_docx(request, pk):
    evento = get_object_or_404(Evento, pk=pk)

    # Congressistas - removendo qualquer limitação de paginação
    inscritos = (
        InscricaoEvento.objects
        .filter(evento=evento)
        .select_related('inscricao')
        .order_by('inscricao__nome')
    )
    
    # Profissionais - removendo qualquer limitação de paginação
    profissionais = (
        ProfissionalEvento.objects
        .filter(evento=evento)
        .select_related('profissional')
        .order_by('profissional__nome')
    )

    # Contadores para UF
    uf_congressistas = Counter()
    uf_profissionais = Counter()

    # Congressistas com UF
    lista_congressistas = []
    for i in inscritos:
        uf = getattr(i.inscricao, 'uf', 'Não informado') or 'Não informado'
        lista_congressistas.append({
            'nome': i.inscricao.nome or 'Sem Cadastro',
            'cpf': i.inscricao.cpf or 'Sem Cadastro',
            'uf': uf
        })
        uf_congressistas[uf] += 1

    # Profissionais com UF
    lista_profissionais = []
    for p in profissionais:
        uf = getattr(p.profissional, 'uf', 'Não informado') or 'Não informado'
        lista_profissionais.append({
            'nome': p.profissional.nome or 'Sem Cadastro',
            'cpf': p.profissional.cpf or 'Sem Cadastro',
            'uf': uf
        })
        uf_profissionais[uf] += 1

    total_congressistas = len(lista_congressistas)
    total_profissionais = len(lista_profissionais)
    total_geral = total_congressistas + total_profissionais

    document = Document()
    document.add_heading(f'Lista - {evento.descricao}', 0)
    document.add_paragraph(f'Gerado em: {now().strftime("%d/%m/%Y %H:%M")}')
    document.add_paragraph('')

    # Congressistas
    document.add_heading('Congressistas', level=1)
    document.add_paragraph(f'Total de congressistas: {total_congressistas}')
    
    for c in lista_congressistas:
        document.add_paragraph(
            f'{c["nome"]} — {c["cpf"]} — UF: {c["uf"]}',
            style='List Number'
        )

    # Resumo UF Congressistas
    document.add_paragraph('')
    document.add_heading('Congressistas por UF', level=2)
    for uf, quantidade in sorted(uf_congressistas.items()):
        document.add_paragraph(f'{uf}: {quantidade} congressista(s)')

    document.add_paragraph('')

    # Profissionais
    document.add_heading('Profissionais', level=1)
    document.add_paragraph(f'Total de profissionais: {total_profissionais}')
    
    for p in lista_profissionais:
        document.add_paragraph(
            f'{p["nome"]} — {p["cpf"]} — UF: {p["uf"]}',
            style='List Number'
        )

    # Resumo UF Profissionais
    document.add_paragraph('')
    document.add_heading('Profissionais por UF', level=2)
    for uf, quantidade in sorted(uf_profissionais.items()):
        document.add_paragraph(f'{uf}: {quantidade} profissional(is)')

    document.add_paragraph('')
    
    # Resumo Geral por UF
    uf_geral = uf_congressistas + uf_profissionais
    document.add_heading('Total Geral por UF', level=1)
    for uf, quantidade in sorted(uf_geral.items()):
        document.add_paragraph(f'{uf}: {quantidade} participante(s)')

    document.add_paragraph('')
    document.add_paragraph(f'Total geral (congressistas + profissionais): {total_geral}')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="evento_{evento.id}_inscritos.docx"'
    document.save(response)
    return response


class ProfissionalRelatorioDocxView(View):
    """ Gera relatório minimalista apenas com dados essenciais """
    
    def get(self, request, *args, **kwargs):
        profissionais = Profissional.objects.all().order_by('nome')
        
        document = Document()
        document.add_heading('Profissionais', 0)
        document.add_paragraph(now().strftime("%d/%m/%Y %H:%M"))
        document.add_paragraph('')

        for profissional in profissionais:
            eventos = ", ".join([e.descricao for e in profissional.eventos.all()])
            
            document.add_paragraph(
                f'{profissional.nome} • {profissional.cpf or "---"} • {eventos or "---"}'
            )

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="profissionais.docx"'
        document.save(response)
        return response


class PlanejamentoRelatorioDocxView(View):
    """ Gera relatório de planejamento com totais e formatação """
    
    def get(self, request, *args, **kwargs):
        planejamentos = Planejamento.objects.all().order_by('descricao')
        total_geral = sum(p.valor_planejado for p in planejamentos if p.valor_planejado)
        
        document = Document()
        document.add_heading('Relatório de Planejamento Financeiro', 0)
        document.add_paragraph(f'Gerado em: {now().strftime("%d/%m/%Y %H:%M")}')
        document.add_paragraph(f'Total de itens: {planejamentos.count()}')
        document.add_paragraph(f'Valor total planejado: R$ {total_geral:,.2f}')
        document.add_paragraph('')

        # Tabela principal
        table = document.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        
        # Cabeçalho
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'DESCRIÇÃO'
        hdr_cells[1].text = 'VALOR PLANEJADO'
        hdr_cells[2].text = 'STATUS'
        
        # Dados
        for planejamento in planejamentos:
            row_cells = table.add_row().cells
            row_cells[0].text = planejamento.descricao or "---"
            
            # Formata valor
            if planejamento.valor_planejado:
                row_cells[1].text = f"R$ {planejamento.valor_planejado:,.2f}"
            else:
                row_cells[1].text = "R$ 0,00"
            
            # Formata status
            status = planejamento.status
            if hasattr(planejamento, 'get_status_display'):
                status = planejamento.get_status_display()
            row_cells[2].text = str(status)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="planejamento_financeiro.docx"'
        document.save(response)
        return response


class PedidosSimplesRelatorioDocxView(View):
    """ Gera relatório minimalista ordenado por nome com resumo """
    
    def get(self, request, *args, **kwargs):
        from django.db.models.functions import Lower
        
        # Ordenar por nome completo em ordem alfabética
        pedidos = PedidoCamisa.objects.all().order_by(Lower('nome_completo'))
        
        # Preparar dados para o resumo
        resumo_por_tipo = {}  # Dicionário para agrupar por tipo de camisa
        
        for pedido in pedidos:
            # Usar o tipo da camisa (choices TIPO_CAMISA) em vez da descrição
            tipo_camisa = pedido.camisa.get_tipo_display() if pedido.camisa else "Sem tipo"
            tamanho = pedido.get_tamanho_display() if hasattr(pedido, 'get_tamanho_display') else pedido.tamanho
            cor = pedido.get_cor_display() if hasattr(pedido, 'get_cor_display') else pedido.cor
            
            # Criar chave combinando tamanho e cor
            chave = f"{tamanho} {cor}"
            
            # Inicializar o dicionário para este tipo de camisa se não existir
            if tipo_camisa not in resumo_por_tipo:
                resumo_por_tipo[tipo_camisa] = {}
            
            # Adicionar ao resumo
            if chave in resumo_por_tipo[tipo_camisa]:
                resumo_por_tipo[tipo_camisa][chave] += 1
            else:
                resumo_por_tipo[tipo_camisa][chave] = 1
        
        document = Document()
        document.add_heading('Pedidos por Ordem Alfabética', 0)
        
        # Data e hora
        data_paragraph = document.add_paragraph()
        data_run = data_paragraph.add_run(now().strftime("%d/%m/%Y %H:%M"))
        data_run.bold = True
        data_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        document.add_paragraph('')
        
        # Adicionar lista de pedidos
        document.add_heading('Lista de Pedidos', level=1)
        
        # Adicionar cabeçalho à lista
        table = document.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Nome'
        hdr_cells[1].text = 'Cidade'
        hdr_cells[2].text = 'Tipo Camisa'
        hdr_cells[3].text = 'Tamanho'
        hdr_cells[4].text = 'Cor'
        hdr_cells[5].text = 'Status'
        
        for pedido in pedidos:
            tamanho = pedido.get_tamanho_display() if hasattr(pedido, 'get_tamanho_display') else pedido.tamanho
            cor = pedido.get_cor_display() if hasattr(pedido, 'get_cor_display') else pedido.cor
            status = pedido.get_status_display() if hasattr(pedido, 'get_status_display') else pedido.status
            tipo_camisa = pedido.camisa.get_tipo_display() if pedido.camisa else "Sem tipo"
            descricao_camisa = pedido.camisa.descricao if pedido.camisa else "Sem descrição"
            
            # Adicionar linha à tabela
            row_cells = table.add_row().cells
            row_cells[0].text = pedido.nome_completo or ''
            row_cells[1].text = pedido.cidade or ''
            row_cells[2].text = f"{tipo_camisa} ({descricao_camisa})"
            row_cells[3].text = tamanho
            row_cells[4].text = cor
            row_cells[5].text = status
        
        document.add_paragraph('')
        
        # Adicionar resumo
        document.add_heading('Resumo por Tipo de Camisa, Tamanho e Cor', level=1)
        
        # Resumo por tipo de camisa (ordenado alfabeticamente)
        for tipo_camisa in sorted(resumo_por_tipo.keys()):
            resumo = resumo_por_tipo[tipo_camisa]
            document.add_heading(f'Camisas {tipo_camisa}', level=2)
            
            if resumo:
                # Ordenar por tamanho e depois por cor
                itens_ordenados = sorted(resumo.items(), key=lambda x: (x[0].split()[0], x[0].split()[1]))
                
                for chave, quantidade in itens_ordenados:
                    partes = chave.split()
                    tamanho = partes[0]
                    cor = ' '.join(partes[1:])
                    
                    p = document.add_paragraph()
                    p.add_run(f'{quantidade} ').bold = True
                    p.add_run(f'{tamanho} {cor}')
            else:
                document.add_paragraph('Nenhum pedido encontrado')
            
            document.add_paragraph('')
        
        # Adicionar totais
        document.add_heading('Totais Gerais', level=2)
        
        total_por_tipo = {}
        for tipo_camisa, resumo in resumo_por_tipo.items():
            total_por_tipo[tipo_camisa] = sum(resumo.values())
        
        total_geral = sum(total_por_tipo.values())
        
        for tipo_camisa, total in sorted(total_por_tipo.items()):
            p_total = document.add_paragraph()
            p_total.add_run(f'Total {tipo_camisa}: ').bold = True
            p_total.add_run(f'{total}')
        
        p_total = document.add_paragraph()
        p_total.add_run('Total Geral: ').bold = True
        p_total.add_run(f'{total_geral}').bold = True

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="pedidos_com_resumo.docx"'
        document.save(response)
        return response


class CaixaCompletoRelatorioDocxView(View):
    """ Gera relatório completo do caixa usando os mesmos cálculos do resumo """
    
    def get(self, request, *args, **kwargs):
        document = Document()
        document.add_heading('Relatório Completo do Caixa', 0)
        document.add_paragraph(f'Gerado em: {now().strftime("%d/%m/%Y %H:%M")}')
        document.add_paragraph('')
        
        # ============ CÁLCULOS (IGUAIS AO RESUMO) ============
        
        # Entradas e saídas
        total_entradas = Entrada.objects.aggregate(total=Sum('valor'))['total'] or 0
        total_saidas = Saida.objects.aggregate(total=Sum('valor'))['total'] or 0

        # Total inscricoes pagas (somando os pagamentos com tipo_modelo='inscricao')
        ct_inscricao = ContentType.objects.get_for_model(Inscricao)
        total_pago_inscricoes = Pagamento.objects.filter(content_type=ct_inscricao).aggregate(total=Sum('valor_pago'))['total'] or 0

        # Total a receber de inscrições (valor_total - valor_pago)
        total_valor_inscricoes = Inscricao.objects.aggregate(total=Sum('valor_total'))['total'] or 0
        total_a_receber_inscricoes = total_valor_inscricoes - total_pago_inscricoes

        # Total camisas - SOMENTE pedidos com status 'pago' ou 'entregue' (JÁ PAGOS)
        total_camisas_pagas = (
            PedidoCamisa.objects
            .filter(Q(status='pago') | Q(status='entregue'))
            .aggregate(
                total=Coalesce(
                    Sum(
                        Case(
                            When(tipo_cliente='equipe', then=Value(0, output_field=DecimalField())),
                            When(tipo_cliente='colaborador', then=F('camisa__valor_compra')),
                            default=F('valor_venda'),
                            output_field=DecimalField()
                        )
                    ),
                    Value(0, output_field=DecimalField())
                )
            )['total']
        )

        # TOTAL DE CAMISAS A RECEBER - pedidos com status 'pendente' ou 'confirmado'
        total_camisas_a_receber = (
            PedidoCamisa.objects
            .filter(Q(status='pendente') | Q(status='confirmado'))
            .aggregate(
                total=Coalesce(
                    Sum(
                        Case(
                            When(tipo_cliente='equipe', then=Value(0, output_field=DecimalField())),
                            When(tipo_cliente='colaborador', then=F('camisa__valor_compra')),
                            default=F('valor_venda'),
                            output_field=DecimalField()
                        )
                    ),
                    Value(0, output_field=DecimalField())
                )
            )['total']
        )

        # Total planejado e total pago em planejamentos
        total_planejamentos = Planejamento.objects.aggregate(total=Sum('valor_planejado'))['total'] or 0
        ct_planejamento = ContentType.objects.get_for_model(Planejamento)
        total_pago_planejamento = Pagamento.objects.filter(content_type=ct_planejamento).aggregate(total=Sum('valor_pago'))['total'] or 0

        # Valor a pagar = planejado - pago
        total_a_pagar = total_planejamentos - total_pago_planejamento

        # Saldo em caixa = entradas + inscrições pagas + camisas pagas - saídas - pagamentos de planejamento
        saldo_caixa = (total_entradas + total_pago_inscricoes + total_camisas_pagas) - (total_saidas + total_pago_planejamento)

        # Cálculo da estimativa futura incluindo camisas a receber
        saldo_futuro_previsto = (saldo_caixa + total_a_receber_inscricoes + total_camisas_a_receber) - total_a_pagar

        # ============ RELATÓRIO ============

        # RESUMO GERAL
        document.add_heading('Resumo Financeiro Geral', level=1)
        
        # SALDO ATUAL
        p = document.add_paragraph()
        p.add_run('Saldo em Caixa: ').bold = True
        p.add_run(f'R$ {saldo_caixa:,.2f}')
        
        document.add_paragraph('')
        
        # ENTRADAS E SAÍDAS
        document.add_heading('Entradas e Saídas', level=2)
        p = document.add_paragraph()
        p.add_run('Total de Entradas: ').bold = True
        p.add_run(f'R$ {total_entradas:,.2f}')
        
        p = document.add_paragraph()
        p.add_run('Total de Saídas: ').bold = True
        p.add_run(f'R$ {total_saidas:,.2f}')
        
        document.add_paragraph('')
        
        # INSCRIÇÕES
        document.add_heading('Inscrições', level=2)
        p = document.add_paragraph()
        p.add_run('Valor Total das Inscrições: ').bold = True
        p.add_run(f'R$ {total_valor_inscricoes:,.2f}')
        
        p = document.add_paragraph()
        p.add_run('Inscrições Pagas: ').bold = True
        p.add_run(f'R$ {total_pago_inscricoes:,.2f}')
        
        p = document.add_paragraph()
        p.add_run('Total a Receber (Inscrições): ').bold = True
        p.add_run(f'R$ {total_a_receber_inscricoes:,.2f}')
        
        document.add_paragraph('')
        
        # CAMISAS
        document.add_heading('Camisas', level=2)
        p = document.add_paragraph()
        p.add_run('Camisas Pagas: ').bold = True
        p.add_run(f'R$ {total_camisas_pagas:,.2f}')
        
        p = document.add_paragraph()
        p.add_run('Total a Receber (Camisas): ').bold = True
        p.add_run(f'R$ {total_camisas_a_receber:,.2f}')
        
        document.add_paragraph('')
        
        # PLANEJAMENTO
        document.add_heading('Planejamento', level=2)
        p = document.add_paragraph()
        p.add_run('Total Planejado: ').bold = True
        p.add_run(f'R$ {total_planejamentos:,.2f}')
        
        p = document.add_paragraph()
        p.add_run('Planejamento Pago: ').bold = True
        p.add_run(f'R$ {total_pago_planejamento:,.2f}')
        
        p = document.add_paragraph()
        p.add_run('Total a Pagar (Planejamento): ').bold = True
        p.add_run(f'R$ {total_a_pagar:,.2f}')
        
        document.add_paragraph('')
        
        # PROJEÇÃO FUTURA
        document.add_heading('Projeção Futura', level=2)
        p = document.add_paragraph()
        p.add_run('Saldo Futuro Previsto: ').bold = True
        p.add_run(f'R$ {saldo_futuro_previsto:,.2f}')

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="relatorio_caixa_simples.docx"'
        document.save(response)
        return response
class RelatorioBaileView(View):
    """ Gera relatório do baile em Word com participantes agrupados por lote """
    
    def get(self, request, baile_id):
        baile = get_object_or_404(BaileAvulso, id=baile_id)
        participantes = (
            ParticipanteBaile.objects.filter(baile=baile)
            .select_related('lote')
            .order_by('lote__descricao', 'nome')
        )
        
        # Agrupar participantes por lote
        participantes_por_lote = {}
        for participante in participantes:
            lote_descricao = participante.lote.descricao if participante.lote else "Sem Lote"
            if lote_descricao not in participantes_por_lote:
                participantes_por_lote[lote_descricao] = []
            participantes_por_lote[lote_descricao].append(participante)
        
        # Calcular totais
        total_arrecadado = sum([
            p.lote.valor_unitario 
            for p in participantes 
            if p.lote and p.lote.valor_unitario
        ])
        lucro = total_arrecadado - (baile.gasto or 0)
        
        # Criar documento Word
        document = Document()
        
        # Cabeçalho
        document.add_heading(f'Relatório do Baile: {baile.nome}', 0)
        
        # Informações do baile
        info_paragraph = document.add_paragraph()
        info_paragraph.add_run('Data: ').bold = True
        info_paragraph.add_run(baile.data.strftime("%d/%m/%Y"))
        
        info_paragraph = document.add_paragraph()
        info_paragraph.add_run('Descrição: ').bold = True
        info_paragraph.add_run(baile.descricao or "---")
        
        info_paragraph = document.add_paragraph()
        info_paragraph.add_run('Gerado em: ').bold = True
        info_paragraph.add_run(now().strftime("%d/%m/%Y %H:%M"))
        
        document.add_paragraph()
        
        # Resumo financeiro
        document.add_heading('Resumo Financeiro', level=1)
        
        resumo_table = document.add_table(rows=4, cols=2)
        resumo_table.style = 'Light Grid Accent 1'
        
        # Dados do resumo
        resumo_data = [
            ('Total de Participantes:', str(participantes.count())),
            ('Total Arrecadado:', f'R$ {total_arrecadado:,.2f}'),
            ('Gastos:', f'R$ {baile.gasto or 0:,.2f}'),
            ('Lucro:', f'R$ {lucro:,.2f}')
        ]
        
        for i, (label, value) in enumerate(resumo_data):
            resumo_table.rows[i].cells[0].text = label
            resumo_table.rows[i].cells[1].text = value
        
        document.add_paragraph()
        
        # Participantes agrupados por lote
        document.add_heading('Participantes por Lote', level=1)
        
        for lote_descricao, participantes_lote in participantes_por_lote.items():
            # Título do lote
            document.add_heading(f'Lote: {lote_descricao}', level=2)
            document.add_paragraph(f'Quantidade: {len(participantes_lote)} participantes')
            
            # Tabela de participantes do lote
            if participantes_lote:
                table = document.add_table(rows=1, cols=3)
                table.style = 'Light Grid Accent 1'
                
                # Cabeçalho
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = '#'
                hdr_cells[1].text = 'Nome'
                hdr_cells[2].text = 'Valor Pago'
                
                # Dados
                for i, participante in enumerate(participantes_lote, 1):
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(i)
                    row_cells[1].text = participante.nome
                    valor = participante.lote.valor_unitario if participante.lote else 0
                    row_cells[2].text = f'R$ {valor:,.2f}'
                
                # Total do lote
                total_lote = sum(p.lote.valor_unitario for p in participantes_lote if p.lote)
                total_row = table.add_row().cells
                total_row[0].text = ''
                total_row[1].text = 'Total do Lote:'
                total_row[2].text = f'R$ {total_lote:,.2f}'
                
                # Formatar célula de total
                for cell in total_row:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
            
            document.add_paragraph()
        
        # Preparar resposta
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="baile_{baile.id}_relatorio.docx"'
        document.save(response)
        
        return response

from django.shortcuts import render
from django.utils.timezone import now
from collections import Counter

class InscricaoRelatorioPreviewView(InscricaoListView):
    """ Pré-visualização do relatório antes de baixar """
    template_name = 'inscricao/inscricao_relatorio_preview.html'
    
    def get(self, request, *args, **kwargs):
        # Remove a paginação para pegar todos os registros
        self.paginate_by = None
        queryset = self.get_queryset()
        
        # Prepara os dados para a pré-visualização
        dados_relatorio = []
        uf_counter = Counter()
        
        for inscricao in queryset:
            categoria = inscricao.categoria.descricao if inscricao.categoria else 'Sem categoria'
            if inscricao.valor_restante_db <= 0:
                status = 'Pago'
            elif inscricao.valor_pago_db > 0:
                status = 'Parcial'
            else:
                status = 'Pendente'
            
            uf = getattr(inscricao, 'uf', 'Não informado') or 'Não informado'
            uf_counter[uf] += 1
            
            dados_relatorio.append({
                'nome': inscricao.nome,
                'cpf': inscricao.cpf,
                'categoria': categoria,
                'status': status,
                'uf': uf,
            })
        
        context = {
            'dados_relatorio': dados_relatorio,
            'total_inscricoes': queryset.count(),
            'data_geracao': now().strftime("%d/%m/%Y %H:%M"),
            'resumo_uf': sorted(uf_counter.items()),
            'filtros_ativos': request.GET.urlencode(),
        }
        
        return render(request, self.template_name, context)